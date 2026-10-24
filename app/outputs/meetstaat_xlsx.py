"""Meetstaat as an Excel workbook: one row per zone x post, numbers as numbers, totals with formulas."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import QuoteResult

HEADERS = [
    "Ruimte/zone",
    "Post",
    "Eenheid",
    "Bruto",
    "Aftrek",
    "Netto",
    "Lagen",
    "Materiaal hoev.",
    "Mat. eenheid",
    "Materiaal €",
    "Uren",
    "Arbeid €",
    "Regel €",
    "Berekening",
]


def build_meetstaat_xlsx(result: QuoteResult, title: str = "Meetstaat") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Meetstaat"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Sector: {result.sector} · prijslijst v{result.config_version} · skill v{result.skill_version}"
    )
    ws["A2"].font = Font(color="666666")

    header_row = 4
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E6E6E6")
        c.alignment = Alignment(wrap_text=True, vertical="top")

    row = header_row + 1
    first = row
    for line in result.lines:
        ws.append(
            [
                line.zone_name,
                line.post_label,
                line.unit,
                line.gross_qty,
                line.deduction_qty,
                line.net_qty,
                line.layers,
                line.material_qty,
                line.material_unit,
                line.material_cost,
                line.hours,
                line.labour_cost,
                line.line_total,
                line.note or "",
            ]
        )
        row += 1
    last = row - 1

    if result.lines:
        ws.cell(row=row, column=1, value="Totaal").font = Font(bold=True)
        for col in (10, 11, 12, 13):
            letter = get_column_letter(col)
            c = ws.cell(row=row, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
            c.font = Font(bold=True)
        row += 2

    t = result.totals
    for label, value in [
        ("Materiaal (excl. marge)", t.material_cost),
        ("Marge op materiaal", t.margin),
        (f"Arbeid ({t.hours_total:g} u)", t.labour_cost),
        ("Subtotaal excl. btw", t.subtotal_excl),
        (f"Btw {t.vat_pct:g} %", t.vat),
        ("Totaal incl. btw", t.total_incl),
    ]:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value).number_format = "#,##0.00"
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Aannames (te bevestigen)").font = Font(bold=True)
    row += 1
    for a in result.assumptions:
        ws.cell(row=row, column=1, value="✔" if a.confirmed else "⚠")
        ws.cell(row=row, column=2, value=a.text)
        row += 1
    for w in result.warnings:
        ws.cell(row=row, column=1, value="!")
        ws.cell(row=row, column=2, value=w)
        row += 1

    widths = [18, 34, 9, 9, 9, 9, 7, 12, 10, 12, 8, 12, 12, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in ws.iter_rows(
        min_row=header_row + 1, max_row=last if result.lines else header_row + 1, min_col=4, max_col=13
    ):
        for c in r:
            if isinstance(c.value, (int, float)):
                c.number_format = "#,##0.00"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
